import re
from pathlib import Path

import requests
import psycopg2
from docx import Document


DB_CONFIG = {
    "dbname": "postgres",
    "user": "postgres",
    "password": "123",
    "host": "localhost",
    "port": "5432",
}

DOCX_FOLDER = Path(r"C:\Users\ignatchenko\Documents\py\pg\docs")
OLLAMA_URL = "http://localhost:11434/api/embeddings"
OLLAMA_MODEL = "nomic-embed-text"


SECTION_TYPES = {
    "Бизнес-процессы": "business_process",
    "Назначение, цель модификации": "purpose",
    "Изменение модели данных": "data_model",
    "Изменение интерфейса пользователей": "interface",
    "Алгоритмы": "algorithm",
    "Настройки": "settings",
    "Настройка": "settings",
    "Отчеты и выходные формы": "reports",
    "Размещение в системе": "placement",
    "Ограничение доступа": "access",
    "Настройка прав доступа": "access",
    "Допущения и ограничения": "limitations",
    "Связанные модификации": "related_modifications",
}

ENTITY_RE = re.compile(
    r"\b[A-Za-zА-Яа-я0-9_]*[A-Za-z]+[A-Za-z0-9_]*\.[A-Za-z0-9_]+"
    r"|\b[A-Z][A-Za-z0-9_]{2,}\b"
    r"|\btsd_[A-Za-z0-9_]+\b"
)

DAX_RE = re.compile(r"DAX-\d+")


def get_embedding(text: str) -> list[float]:
    response = requests.post(
        OLLAMA_URL,
        json={"model": OLLAMA_MODEL, "prompt": text},
        timeout=120,
    )
    response.raise_for_status()
    return response.json()["embedding"]


def read_docx_text(path: Path) -> str:
    doc = Document(path)
    parts = []

    for p in doc.paragraphs:
        text = p.text.strip()
        if text:
            parts.append(text)

    # Таблицы тоже читаем, потому что в твоих ФД там лежат поля/enum/история.
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            line = " | ".join([c for c in cells if c])
            if line:
                parts.append(line)

    return "\n".join(parts)


def parse_doc_metadata(text: str, file_name: str) -> dict:
    dax = DAX_RE.search(file_name) or DAX_RE.search(text)

    title = file_name.rsplit(".", 1)[0]
    first_lines = [line.strip() for line in text.splitlines() if line.strip()]
    if first_lines:
        title = first_lines[0]

    version = None
    date = None
    author = None

    version_match = re.search(r"\b0\.\d+\b", text)
    if version_match:
        version = version_match.group(0)

    date_match = re.search(r"\b\d{2}\.\d{2}\.\d{4}\b", text)
    if date_match:
        d, m, y = date_match.group(0).split(".")
        date = f"{y}-{m}-{d}"

    author_match = re.search(r"\b(Юрьева Ульяна|Игнатченко Евгений|Matveev Dmitriy)\b", text)
    if author_match:
        author = author_match.group(1)

    business_process = extract_section_text(text, "Бизнес-процессы", max_chars=500)
    purpose = extract_section_text(text, "Назначение, цель модификации", max_chars=1000)

    return {
        "dax_code": dax.group(0) if dax else None,
        "title": title,
        "source_file": file_name,
        "version": version,
        "business_process": business_process,
        "purpose": purpose,
        "author": author,
        "document_date": date,
        "source_type": "docx",
    }


def extract_section_text(text: str, section_title: str, max_chars: int = 1000) -> str | None:
    lines = text.splitlines()
    capture = False
    result = []

    known_titles = set(SECTION_TYPES.keys())

    for line in lines:
        clean = line.strip()
        if clean == section_title:
            capture = True
            continue

        if capture and clean in known_titles:
            break

        if capture and clean:
            result.append(clean)

    value = "\n".join(result).strip()
    return value[:max_chars] if value else None


def split_into_sections(text: str) -> list[dict]:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    sections = []
    current_title = "Общее"
    current_type = "general"
    current_lines = []

    def flush():
        if current_lines:
            sections.append({
                "section_title": current_title,
                "chunk_type": current_type,
                "text": "\n".join(current_lines).strip(),
            })

    for line in lines:
        if line in SECTION_TYPES:
            flush()
            current_title = line
            current_type = SECTION_TYPES[line]
            current_lines = []
        else:
            current_lines.append(line)

    flush()
    return sections


def chunk_text(text: str, max_chars: int = 1800, overlap: int = 250) -> list[str]:
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) <= max_chars:
        return [text]

    chunks = []
    start = 0

    while start < len(text):
        end = start + max_chars
        piece = text[start:end]

        # Пытаемся резать по концу абзаца/предложения.
        cut = max(piece.rfind("\n"), piece.rfind(". "), piece.rfind("; "))
        if cut > 500:
            piece = piece[:cut + 1]
            end = start + cut + 1

        chunks.append(piece.strip())
        start = max(end - overlap, end)

    return [c for c in chunks if c]


def extract_entities(text: str) -> list[str]:
    found = set()

    for match in ENTITY_RE.findall(text):
        value = match.strip(".,;:()[]«»")
        if len(value) >= 3:
            found.add(value)

    # Дополнительно ловим частые бизнес-сущности из твоих ФД.
    keywords = [
        "WMS_IsOversizedItemIM",
        "LFL_SCSPackTask",
        "WMSPickingRoute",
        "WMSOrderTrans",
        "PickingLineBuffer",
        "SalesTable",
        "InventTable",
        "WMSStoreArea",
        "WMS_TSDTaskType",
        "WMS_OperationType",
        "tsd_Setup",
        "tsd_AutoTaskTable",
        "tsd_WMSLocation",
    ]

    for k in keywords:
        if k in text:
            found.add(k)

    return sorted(found)


def insert_document(cur, meta: dict) -> int:
    cur.execute(
        """
        INSERT INTO ai.fd_documents
            (title, source_file, dax_code, version, business_process, purpose, author, document_date, source_type)
        VALUES
            (%s, %s, %s, %s, %s, %s, %s, %s, %s)
        RETURNING id
        """,
        (
            meta["title"],
            meta["source_file"],
            meta["dax_code"],
            meta["version"],
            meta["business_process"],
            meta["purpose"],
            meta["author"],
            meta["document_date"],
            meta["source_type"],
        ),
    )
    return cur.fetchone()[0]


def insert_entity(cur, name: str) -> int:
    cur.execute(
        """
        INSERT INTO ai.fd_entities (name)
        VALUES (%s)
        ON CONFLICT (name) DO UPDATE SET name = EXCLUDED.name
        RETURNING id
        """,
        (name,),
    )
    return cur.fetchone()[0]


def insert_chunk(cur, document_id: int, chunk: str, section_title: str, chunk_type: str, entities: list[str]) -> int:
    embedding = get_embedding(chunk)

    cur.execute(
        """
        INSERT INTO ai.fd_chunks
            (document_id, chunk_text, embedding, section_title, chunk_type, business_terms, jira_keys, process_type)
        VALUES
            (%s, %s, %s, %s, %s, %s, %s, %s)
        RETURNING id
        """,
        (
            document_id,
            chunk,
            embedding,
            section_title,
            chunk_type,
            [],
            DAX_RE.findall(chunk),
            guess_process_type(chunk),
        ),
    )
    chunk_id = cur.fetchone()[0]

    for entity in entities:
        entity_id = insert_entity(cur, entity)
        cur.execute(
            """
            INSERT INTO ai.fd_chunk_entities (chunk_id, entity_id)
            VALUES (%s, %s)
            ON CONFLICT DO NOTHING
            """,
            (chunk_id, entity_id),
        )

    return chunk_id


def guess_process_type(text: str) -> str | None:
    lower = text.lower()

    if "упаков" in lower:
        return "packing"
    if "тсд" in lower or "tsd_" in text:
        return "tsd"
    if "сортиров" in lower or "scs" in lower:
        return "sorting"
    if "комплектац" in lower or "сборк" in lower:
        return "picking"
    if "номенклатур" in lower or "inventtable" in lower:
        return "item_master"

    return None


def insert_links_after_load(cur):
    cur.execute("SELECT id, dax_code FROM ai.fd_documents WHERE dax_code IS NOT NULL")
    docs = {dax: doc_id for doc_id, dax in cur.fetchall()}

    cur.execute("""
        SELECT c.document_id, c.chunk_text
        FROM ai.fd_chunks c
        WHERE c.chunk_type = 'related_modifications'
    """)

    for source_document_id, text in cur.fetchall():
        for dax in set(DAX_RE.findall(text)):
            target_id = docs.get(dax)

            if target_id and target_id != source_document_id:
                cur.execute(
                    """
                    INSERT INTO ai.fd_document_links
                        (source_document_id, target_document_id, link_type)
                    VALUES
                        (%s, %s, %s)
                    ON CONFLICT DO NOTHING
                    """,
                    (source_document_id, target_id, "related"),
                )



def main():
    files = sorted({p.resolve() for p in DOCX_FOLDER.glob("*") if p.suffix.lower() == ".docx"})

    if not files:
        print(f"Не найдены DOCX в папке: {DOCX_FOLDER}")
        return

    conn = psycopg2.connect(**DB_CONFIG)
    cur = conn.cursor()

    for path in files:
        print(f"Загружаю: {path.name}")

        text = read_docx_text(path)
        meta = parse_doc_metadata(text, path.name)
        document_id = insert_document(cur, meta)

        sections = split_into_sections(text)

        chunk_count = 0
        for section in sections:
            for chunk in chunk_text(section["text"]):
                entities = extract_entities(chunk)
                insert_chunk(
                    cur=cur,
                    document_id=document_id,
                    chunk=chunk,
                    section_title=section["section_title"],
                    chunk_type=section["chunk_type"],
                    entities=entities,
                )
                chunk_count += 1

        print(f"  document_id={document_id}, chunks={chunk_count}")

    insert_links_after_load(cur)

    conn.commit()
    cur.close()
    conn.close()

    print("Готово: документы загружены")


if __name__ == "__main__":
    main()