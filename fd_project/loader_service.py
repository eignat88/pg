from __future__ import annotations

import re
from pathlib import Path

from docx import Document

from .config import AppConfig
from .db import get_connection
from .embeddings import EmbeddingClient
from .repositories import FDRepository

SECTION_TYPES = {
    "Бизнес-процессы": "business_process",
    "Назначение, цель модификации": "purpose",
    "Изменение модели данных": "data_model",
    "Изменение интерфейса пользователей": "interface",
    "Алгоритмы": "algorithm",
    "Настройки": "settings",
    "Настройка": "settings",
    "Отчеты и выходные формы": "reports",
    "Размещение в системе": "placement",
    "Ограничение доступа": "access",
    "Настройка прав доступа": "access",
    "Допущения и ограничения": "limitations",
    "Связанные модификации": "related_modifications",
}

ENTITY_RE = re.compile(
    r"\b[A-Za-zА-Яа-я0-9_]*[A-Za-z]+[A-Za-z0-9_]*\.[A-Za-z0-9_]+"
    r"|\b[A-Z][A-Za-z0-9_]{2,}\b"
    r"|\btsd_[A-Za-z0-9_]+\b"
)
DAX_RE = re.compile(r"DAX-\d+")


def read_docx_text(path: Path) -> str:
    doc = Document(path)
    parts: list[str] = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if text:
            parts.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            line = " | ".join([c for c in cells if c])
            if line:
                parts.append(line)
    return "\n".join(parts)


def extract_section_text(text: str, section_title: str, max_chars: int = 1000) -> str | None:
    lines = text.splitlines()
    capture = False
    result = []
    known_titles = set(SECTION_TYPES.keys())

    for line in lines:
        clean = line.strip()
        if clean == section_title:
            capture = True
            continue
        if capture and clean in known_titles:
            break
        if capture and clean:
            result.append(clean)

    value = "\n".join(result).strip()
    return value[:max_chars] if value else None


def parse_doc_metadata(text: str, file_name: str) -> dict:
    dax = DAX_RE.search(file_name) or DAX_RE.search(text)
    title = file_name.rsplit(".", 1)[0]
    first_lines = [line.strip() for line in text.splitlines() if line.strip()]
    if first_lines:
        title = first_lines[0]

    version = re.search(r"\b0\.\d+\b", text)
    date_match = re.search(r"\b\d{2}\.\d{2}\.\d{4}\b", text)
    author_match = re.search(r"\b(Юрьева Ульяна|Игнатченко Евгений|Matveev Dmitriy)\b", text)

    doc_date = None
    if date_match:
        d, m, y = date_match.group(0).split(".")
        doc_date = f"{y}-{m}-{d}"

    return {
        "dax_code": dax.group(0) if dax else None,
        "title": title,
        "source_file": file_name,
        "version": version.group(0) if version else None,
        "business_process": extract_section_text(text, "Бизнес-процессы", max_chars=500),
        "purpose": extract_section_text(text, "Назначение, цель модификации", max_chars=1000),
        "author": author_match.group(1) if author_match else None,
        "document_date": doc_date,
        "source_type": "docx",
    }


def split_into_sections(text: str) -> list[dict]:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    sections = []
    current_title, current_type, current_lines = "Общее", "general", []

    def flush() -> None:
        if current_lines:
            sections.append({"section_title": current_title, "chunk_type": current_type, "text": "\n".join(current_lines).strip()})

    for line in lines:
        if line in SECTION_TYPES:
            flush()
            current_title, current_type, current_lines = line, SECTION_TYPES[line], []
        else:
            current_lines.append(line)
    flush()
    return sections


def chunk_text(text: str, max_chars: int = 1800, overlap: int = 250) -> list[str]:
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) <= max_chars:
        return [text]

    chunks: list[str] = []
    start = 0
    while start < len(text):
        end = start + max_chars
        piece = text[start:end]
        cut = max(piece.rfind("\n"), piece.rfind(". "), piece.rfind("; "))
        if cut > 500:
            piece, end = piece[: cut + 1], start + cut + 1
        chunks.append(piece.strip())
        start = max(end - overlap, end)
    return [c for c in chunks if c]


def extract_entities(text: str) -> list[str]:
    found = {m.strip(".,;:()[]«»") for m in ENTITY_RE.findall(text) if len(m.strip(".,;:()[]«»")) >= 3}
    return sorted(found)


def guess_process_type(text: str) -> str | None:
    lower = text.lower()
    if "упаков" in lower:
        return "packing"
    if "тсд" in lower or "tsd_" in text:
        return "tsd"
    if "сортиров" in lower or "scs" in lower:
        return "sorting"
    if "комплектац" in lower or "сборк" in lower:
        return "picking"
    if "номенклатур" in lower or "inventtable" in lower:
        return "item_master"
    return None


def rebuild_document_links(repo: FDRepository) -> None:
    repo.cur.execute("DELETE FROM ai.fd_document_links")
    repo.cur.execute("SELECT id, dax_code FROM ai.fd_documents WHERE dax_code IS NOT NULL")
    docs = {dax: doc_id for doc_id, dax in repo.cur.fetchall()}
    repo.cur.execute("SELECT document_id, chunk_text FROM ai.fd_chunks WHERE chunk_type = 'related_modifications'")

    for source_document_id, text in repo.cur.fetchall():
        for dax in set(DAX_RE.findall(text)):
            target_id = docs.get(dax)
            if target_id and target_id != source_document_id:
                repo.insert_document_link(source_document_id, target_id, "related")


def load_docx_batch(config: AppConfig | None = None, replace_existing: bool = True) -> None:
    config = config or AppConfig()
    files = sorted({p.resolve() for p in config.docx_folder.glob("*") if p.suffix.lower() == ".docx"})
    if not files:
        print(f"Не найдены DOCX в папке: {config.docx_folder}")
        return

    embedding_client = EmbeddingClient(config.ollama_url, config.ollama_model)
    with get_connection(config) as conn:
        cur = conn.cursor()
        repo = FDRepository(cur)

        for path in files:
            print(f"Загружаю: {path.name}")
            if replace_existing:
                old_id = repo.find_document_id_by_source_file(path.name)
                if old_id:
                    repo.delete_document_graph(old_id)

            text = read_docx_text(path)
            document_id = repo.insert_document(parse_doc_metadata(text, path.name))

            chunk_count = 0
            for section in split_into_sections(text):
                for chunk in chunk_text(section["text"]):
                    chunk_id = repo.insert_chunk(
                        {
                            "document_id": document_id,
                            "chunk_text": chunk,
                            "embedding": embedding_client.get_embedding(chunk),
                            "section_title": section["section_title"],
                            "chunk_type": section["chunk_type"],
                            "business_terms": [],
                            "jira_keys": DAX_RE.findall(chunk),
                            "process_type": guess_process_type(chunk),
                        }
                    )
                    for entity_name in extract_entities(chunk):
                        repo.link_chunk_entity(chunk_id, repo.insert_entity(entity_name))
                    chunk_count += 1

            print(f"  document_id={document_id}, chunks={chunk_count}")

        rebuild_document_links(repo)
        conn.commit()

        stats = repo.get_table_stats()
        print(
            "Готово. Текущие размеры (строки): "
            f"fd_chunks={stats.fd_chunks}, fd_documents={stats.fd_documents}, "
            f"fd_document_links={stats.fd_document_links}, fd_entities={stats.fd_entities}, "
            f"fd_chunk_entities={stats.fd_chunk_entities}"
        )
